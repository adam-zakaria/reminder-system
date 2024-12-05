import openpyxl
from openpyxl import Workbook
import os
import json

def init_excel(file_name='conversation_store.xlsx'):
    try:
        # Check if file already exists
        if not os.path.exists(file_name):
            # Create a new workbook and sheets if the file doesn't exist
            wb = Workbook()
            wb.create_sheet('Sessions')

            if 'Sheet' in wb.sheetnames:
                wb.remove(wb['Sheet'])

            # Add header row to 'Sessions' sheet
            sheet = wb['Sessions']
            sheet.append(['Session ID', 'Conversation', 'Summarization', 'Generated Code', 'Analysed Data'])

            wb.save(file_name)
            print(f"Excel file '{file_name}' created successfully.")
        else:
            print(f"Excel file '{file_name}' already exists.")

    except Exception as e:
        print(f"Error initializing Excel file: {e}")
        
def serialize_data(data):
    # Convert complex data (like dicts) to a JSON string
    if isinstance(data, (dict, list)):
        return json.dumps(data, indent=2)
    return data

def update_session_data(session_id, conversation=None, summarization=None, code_generation=None, analysed_data=None, file_name='conversation_store.xlsx'):
    try:
        # Ensure the file exists before accessing
        if not os.path.exists(file_name):
            init_excel(file_name)
        
        wb = openpyxl.load_workbook(file_name)
        sheet = wb['Sessions']
        rows = list(sheet.iter_rows(values_only=True))
        
        # Serialize data before writing to Excel
        conversation = serialize_data(conversation)
        summarization = serialize_data(summarization)
        code_generation = serialize_data(code_generation)
        analysed_data = serialize_data(analysed_data)

        # Find existing session row
        row_index = None
        for i, row in enumerate(rows):
            if row[0] == session_id:
                row_index = i + 1  # 1-based index for Excel
                break

        if row_index is None:
            # Add a new row for a new session
            sheet.append([session_id, conversation, summarization, code_generation, analysed_data])
        else:
            # Update the existing row for the session
            if conversation:
                sheet.cell(row=row_index, column=2, value=conversation)
            if summarization:
                sheet.cell(row=row_index, column=3, value=summarization)
            if code_generation:
                sheet.cell(row=row_index, column=4, value=code_generation)
            if analysed_data:
                sheet.cell(row=row_index, column=5, value=analysed_data)

        wb.save(file_name)
        print(f"Session data for {session_id} updated successfully.")

    except Exception as e:
        print(f"Error updating session data in Excel: {e}")
